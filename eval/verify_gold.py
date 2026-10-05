#!/usr/bin/env python3
"""Independent verifier for golden_traces.json.

Re-derives every machine-checkable evidence item from the raw dataset, so the gold set can be audited
without trusting whoever wrote it. Exit code 0 only if every check passes.

Usage:
  python verify_gold.py --traces golden_traces.json \
      --dataset-root task-data/aegis-dataset/aegis-dataset \
      --eval-pdf task-data/evaluation-questions.pdf

Needs: poppler-utils (pdftotext), pdfplumber, openpyxl, python-docx, python-pptx, pydantic>=2.
"""
from __future__ import annotations
import argparse, html, json, re, subprocess, sys
from enum import Enum
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal

import openpyxl, pdfplumber
from docx import Document
from pptx import Presentation
from pydantic import BaseModel, ConfigDict, model_validator


# ----------------------------------------------------------------------------- schema
class Status(str, Enum):
    ANSWERED = "ANSWERED"; ANSWERED_SCOPED = "ANSWERED_SCOPED"; CONFLICTED = "CONFLICTED"
    PARTIAL = "PARTIAL"; UNANSWERABLE = "UNANSWERABLE"

class Kind(str, Enum):
    stated = "stated"; inferred = "inferred"; structural = "structural"

class Claim(BaseModel):
    model_config = ConfigDict(extra="forbid")
    key: str; entity: str; predicate: str; value: Any; scope: dict; assertion_kind: Kind; unit: str | None = None

class Slot(BaseModel):
    model_config = ConfigDict(extra="forbid")
    entity: str; predicate: str; required: bool; allowed_assertion_kinds: list[Kind]

class Evidence(BaseModel):
    model_config = ConfigDict(extra="forbid")
    role: str; document: str; locator: dict; check: dict; note: str | None = None

class Trace(BaseModel):
    model_config = ConfigDict(extra="forbid")
    question_id: str; question_text: str; intent: str; target_entities: list[str]
    required_slots: list[Slot]; scope: dict; expected_status: Status; trap_types: list[str]
    expected_claims: list[Claim]; expected_evidence: list[Evidence]; required_values: list[dict]
    gaps: list[str]; must_mention: list[str]; must_not_assert: list[str]
    critical_errors: list[str]; major_errors: list[str]; consistency_group: str | None
    gold_answer_summary: str; reviewer_notes: str | None; reviewer_decision_ref: str | None

    @model_validator(mode="after")
    def _semantics(self):
        s, ev = self.expected_status, self.expected_evidence
        if s == Status.UNANSWERABLE:
            assert not self.expected_claims, "UNANSWERABLE must have no expected claims"
            assert any(e.role == "checked_no_fact" for e in ev), "UNANSWERABLE needs checked_no_fact evidence"
            assert self.gaps, "UNANSWERABLE needs gaps"
        if s == Status.PARTIAL:
            assert self.expected_claims and self.gaps, "PARTIAL needs both found claims and gaps"
        if s in (Status.ANSWERED, Status.ANSWERED_SCOPED):
            assert any(e.role == "primary" for e in ev), "answerable traces need primary evidence"
            assert self.expected_claims, "answerable traces need expected claims"
        assert self.critical_errors, "every trace must define critical errors"
        return self

class Doc(BaseModel):
    model_config = ConfigDict(extra="allow")
    traces: list[Trace]; reviewer_decisions: list[dict]


# ----------------------------------------------------------------------------- text helpers
_Q = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u00a0": " "})
def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.translate(_Q)).strip()

class Corpus:
    def __init__(self, root: Path): self.root = root

    @lru_cache(maxsize=None)
    def pdf_page(self, rel: str, page: int) -> str:
        out = subprocess.run(["pdftotext", "-layout", "-f", str(page), "-l", str(page), str(self.root / rel), "-"],
                             capture_output=True, text=True, check=True).stdout
        return norm(out)

    @lru_cache(maxsize=None)
    def pdf_pages(self, rel: str) -> int:
        with pdfplumber.open(self.root / rel) as p: return len(p.pages)

    @lru_cache(maxsize=None)
    def xlsx(self, rel: str) -> dict[str, list[dict[str, str]]]:
        wb = openpyxl.load_workbook(self.root / rel, data_only=True)
        res = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            hdr = [str(h) for h in rows[0]]
            res[ws.title] = [{h: ("" if v is None else str(v)) for h, v in zip(hdr, r)} for r in rows[1:]]
        return res

    @lru_cache(maxsize=None)
    def json_(self, rel: str) -> Any: return json.loads((self.root / rel).read_text())

    @lru_cache(maxsize=None)
    def docx(self, rel: str) -> str:
        d = Document(self.root / rel)
        parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
        return norm(" ".join(parts))

    @lru_cache(maxsize=None)
    def pptx_slide(self, rel: str, n: int) -> str:
        sl = Presentation(self.root / rel).slides[n - 1]; parts = []
        for sh in sl.shapes:
            if sh.has_text_frame: parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [c.text for r in sh.table.rows for c in r.cells]
        return norm(" ".join(parts))

    @lru_cache(maxsize=None)
    def html(self, rel: str) -> str:
        t = (self.root / rel).read_text()
        t = re.sub(r"<(script|style).*?</\1>", "", t, flags=re.S)
        return norm(html.unescape(re.sub(r"<[^>]+>", " ", t)))

    def all_text(self) -> dict[str, str]:
        """Every text-bearing source (images and the scan have no reliable text layer)."""
        out = {}
        for p in sorted(self.root.rglob("*")):
            rel = str(p.relative_to(self.root))
            if p.suffix == ".pdf": out[rel] = " ".join(self.pdf_page(rel, i + 1) for i in range(self.pdf_pages(rel)))
            elif p.suffix == ".xlsx": out[rel] = norm(" ".join(v for rows in self.xlsx(rel).values() for r in rows for v in r.values()))
            elif p.suffix == ".docx": out[rel] = self.docx(rel)
            elif p.suffix == ".pptx": out[rel] = norm(" ".join(self.pptx_slide(rel, i + 1) for i in range(len(Presentation(p).slides))))
            elif p.suffix == ".html": out[rel] = self.html(rel)
            elif p.suffix == ".json": out[rel] = norm(p.read_text())
        return out


# ----------------------------------------------------------------------------- checks
MISSING = object()
def dig(obj: Any, path: str) -> Any:
    for k in path.split("."):
        if not isinstance(obj, dict) or k not in obj: return MISSING
        obj = obj[k]
    return obj

def geometry_edges(root: Path, rel: str, eps: float = 3.0):
    """Ground-truth graph from vector geometry: boxes = filled curves; edges = lines whose endpoints lie on box edges."""
    with pdfplumber.open(root / rel) as pdf:
        pg = pdf.pages[0]
        words = pg.extract_words()
        boxes = []
        for c in pg.curves:
            inside = [w["text"] for w in words if c["x0"] - 1 <= w["x0"] and w["x1"] <= c["x1"] + 1 and c["top"] - 1 <= w["top"] and w["bottom"] <= c["bottom"] + 1]
            boxes.append((c["x0"], c["top"], c["x1"], c["bottom"], " ".join(inside)))
        def which(x, y):
            for x0, t, x1, b, label in boxes:
                if x0 - eps <= x <= x1 + eps and t - eps <= y <= b + eps: return label
            return None
        def colour(c):
            r, g, b = (round(v, 2) for v in c)
            return "purple" if (r, g, b) == (0.54, 0.29, 0.61) else "tan" if (r, g, b) == (0.69, 0.54, 0.31) else "grey"
        edges = []
        for l in pg.lines:
            (x0, y0), (x1, y1) = l["pts"][0], l["pts"][-1]
            edges.append((colour(l["stroking_color"]), which(x0, y0), which(x1, y1)))
        return boxes, edges

def short(label: str | None) -> str | None:
    if label is None: return None
    for k in ("PLC-03", "PS-04A", "IV-21", "Hydraulic Power Unit", "Main Reservoir", "Auxiliary", "Press Circuit"):
        if k in label: return {"Hydraulic Power Unit": "HPU", "Auxiliary": "Auxiliary Reservoir"}.get(k, k)
    return label

def run_check(ev: Evidence, C: Corpus, root: Path) -> tuple[str, str]:
    """returns (PASS|FAIL|MANUAL, detail)"""
    ck, t = ev.check, ev.check["type"]
    if t == "manual_visual": return "MANUAL", ck["observed"][:70]
    if t == "pdf_text":
        txt = C.pdf_page(ev.document, ck["page"])
        bad = [q for q in ck["quotes"] if norm(q) not in txt]
        return ("PASS", f"{len(ck['quotes'])} quote(s)") if not bad else ("FAIL", f"not on p{ck['page']}: {bad}")
    if t == "xlsx_cell":
        rows = C.xlsx(ev.document).get(ck["sheet"], [])
        row = next((r for r in rows if r.get(ck["key_col"]) == ck["key"]), None)
        if row is None: return "FAIL", f"row {ck['key']!r} not found"
        v = row.get(ck["col"], "")
        if "equals" in ck and v != ck["equals"]: return "FAIL", f"{ck['col']}={v!r} != {ck['equals']!r}"
        if "contains" in ck and ck["contains"] not in v: return "FAIL", f"{v!r} lacks {ck['contains']!r}"
        return "PASS", f"{ck['key']}.{ck['col']}"
    if t == "json_path":
        v = dig(C.json_(ev.document), ck["path"])
        if v is MISSING: return "FAIL", f"path {ck['path']} missing"
        return ("PASS", f"{ck['path']}={v!r}") if v == ck["equals"] else ("FAIL", f"{ck['path']}={v!r} != {ck['equals']!r}")
    if t in ("docx_text", "pptx_text", "html_text"):
        txt = C.docx(ev.document) if t == "docx_text" else C.pptx_slide(ev.document, ck["slide"]) if t == "pptx_text" else C.html(ev.document)
        bad = [q for q in ck["quotes"] if norm(q) not in txt]
        return ("PASS", f"{len(ck['quotes'])} quote(s)") if not bad else ("FAIL", f"missing: {bad}")
    if t == "geometry_hydraulic":
        boxes, edges = geometry_edges(root, ev.document)
        purple = {frozenset((short(a), short(b))) for c, a, b in edges if c == "purple"}
        want = {frozenset(p) for p in ck["purple_edges"]}
        grey_ctrl = [(short(a), short(b)) for c, a, b in edges if c == "grey" and "PLC-03" in (short(a), short(b))]
        info = f"purple={sorted(map(sorted, purple))} | grey lines touching PLC-03={grey_ctrl}"
        return ("PASS", info) if purple == want else ("FAIL", f"expected {sorted(map(sorted, want))}; {info}")
    if t == "absence":
        hits = []
        for rel, txt in C.all_text().items():
            if rel in ck["allowed_in"]: continue
            for pat in ck["patterns"]:
                m = re.search(pat, txt, flags=re.I)
                if m: hits.append(f"{rel}: {pat!r} -> ...{txt[max(0, m.start() - 30):m.end() + 30]}...")
        return ("PASS", f"{len(ck['patterns'])} pattern(s) absent outside allowed docs") if not hits else ("FAIL", "; ".join(hits))
    return "FAIL", f"unknown check type {t}"


def eval_questions(pdf_path: Path) -> dict[str, str]:
    txt = subprocess.run(["pdftotext", "-layout", str(pdf_path), "-"], capture_output=True, text=True, check=True).stdout
    qs, cur = {}, None
    for line in txt.splitlines():
        m = re.match(r"\s*(\d+)\.\s+(.*)", line)
        if m: cur = f"Q{m.group(1)}"; qs[cur] = m.group(2).strip()
        elif cur and line.strip() and "Evaluation Questions" not in line: qs[cur] += " " + line.strip()
    return {k: norm(v) for k, v in qs.items()}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--traces", default="golden_traces.json"); ap.add_argument("--dataset-root", required=True)
    ap.add_argument("--eval-pdf"); ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(); root = Path(a.dataset_root); C = Corpus(root)
    fails = 0

    raw = json.loads(Path(a.traces).read_text())
    try: doc = Doc.model_validate(raw)
    except Exception as e: print("SCHEMA FAIL:", e); return 2
    ids = {d["id"] for d in doc.reviewer_decisions}
    for t in doc.traces:
        if t.reviewer_decision_ref and t.reviewer_decision_ref not in ids: print(f"{t.question_id}: dangling decision ref"); fails += 1
    groups: dict[str, list[Trace]] = {}
    for t in doc.traces:
        if t.consistency_group and t.required_values: groups.setdefault(t.consistency_group, []).append(t)
    for g, ts in groups.items():
        vals = {json.dumps(sorted(json.dumps(v, sort_keys=True) for v in t.required_values)) for t in ts}
        if len(vals) > 1: print(f"consistency group {g}: members disagree on required_values"); fails += 1
    print(f"schema OK: {len(doc.traces)} traces, statuses:", {s.value: sum(t.expected_status == s for t in doc.traces) for s in Status})

    if a.eval_pdf:
        gold = eval_questions(Path(a.eval_pdf))
        for t in doc.traces:
            ok = gold.get(t.question_id) == norm(t.question_text)
            if not ok: print(f"{t.question_id}: question text NOT verbatim\n  gold: {gold.get(t.question_id)}\n  ours: {norm(t.question_text)}"); fails += 1
        print(f"question text: {sum(gold.get(t.question_id) == norm(t.question_text) for t in doc.traces)}/{len(doc.traces)} verbatim vs evaluation PDF")

    tally = {"PASS": 0, "FAIL": 0, "MANUAL": 0}; manual = []
    for t in doc.traces:
        res = [(e, *run_check(e, C, root)) for e in t.expected_evidence]
        for e, st, d in res:
            tally[st] += 1
            if st == "FAIL": fails += 1; print(f"  FAIL {t.question_id} [{e.role}] {e.document}: {d}")
            elif st == "MANUAL": manual.append((t.question_id, e.document, d))
            elif a.verbose: print(f"  pass {t.question_id} [{e.role}] {e.document}: {d}")
        print(f"{t.question_id:>3} {t.expected_status.value:<16} " + " ".join({"PASS": ".", "FAIL": "X", "MANUAL": "m"}[s] for _, s, _ in res))
    print(f"\nevidence checks: {tally}")
    print("\nNEEDS A HUMAN EYE (image/scan content; OCR is not trustworthy here):")
    for q, d, o in manual: print(f"  {q}: {d} :: {o}")
    print("\nRESULT:", "OK" if fails == 0 else f"{fails} failure(s)")
    return 0 if fails == 0 else 1

if __name__ == "__main__":
    sys.exit(main())
